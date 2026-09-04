# -*- coding: utf-8 -*-
"""报价助手路由：客户画像 CRUD、报价Excel导入与统计、附件管理、AI提示词、导出"""
import json
import io
import base64
from datetime import datetime

from db import query_all, query_one, tx, log

try:
    import openpyxl
except ImportError:
    openpyxl = None


def routes():
    return [
        # 客户画像 CRUD
        ('GET', r'/api/customer-profiles/(?P<customer_id>\d+)', h_get_profile),
        ('POST', r'/api/customer-profiles/(?P<customer_id>\d+)', h_save_profile),
        # 报价 Excel 导入与统计
        ('POST', r'/api/customer-profiles/(?P<customer_id>\d+)/import-excel', h_import_excel),
        ('GET', r'/api/customer-profiles/(?P<customer_id>\d+)/stats', h_get_stats),
        ('GET', r'/api/customer-profiles/(?P<customer_id>\d+)/excel-list', h_excel_list),
        ('DELETE', r'/api/customer-quote-excels/(?P<id>\d+)', h_delete_excel),
        # 报价历史同步统计
        ('GET', r'/api/customer-profiles/(?P<customer_id>\d+)/quotation-stats', h_quotation_stats),
        # 附件管理
        ('POST', r'/api/customer-profiles/(?P<customer_id>\d+)/attachments', h_upload_attachment),
        ('GET', r'/api/customer-profiles/(?P<customer_id>\d+)/attachments', h_list_attachments),
        ('DELETE', r'/api/customer-attachments/(?P<id>\d+)', h_delete_attachment),
        # AI 提示词模板
        ('GET', r'/api/ai-prompts', h_get_prompts),
        # 导出
        ('GET', r'/api/customer-profiles/(?P<customer_id>\d+)/export/(?P<fmt>md|json)', h_export_profile),
    ]


# ==================== 客户画像 CRUD ====================

def h_get_profile(p, q, b, http):
    cid = int(p['customer_id'])
    cust = query_one('SELECT * FROM customers WHERE id=?', (cid,))
    if not cust:
        return http.send_json({'error': '客户不存在'}, 404)
    pf = query_one('SELECT * FROM customer_profiles WHERE customer_id=?', (cid,))
    if not pf:
        # 自动从 customers 表预填基础字段
        pf = {
            'id': None, 'customer_id': cid,
            'company_scale': '', 'main_products': '', 'website': '',
            'comm_best_time': '', 'comm_style': '', 'is_urgent_order': 0,
            'own_forwarder': cust.get('incoterm', '') or '',
            'custom_level': '', 'certification': '',
            'trade_terms_detail': cust.get('incoterm', '') or '',
            'risk_notes': '', 'profile_notes': cust.get('notes', '') or '',
            'quote_summary': '', 'website_checked': 0, 'updated_at': ''
        }
    pf['customer'] = cust
    return pf


def h_save_profile(p, q, b, http):
    cid = int(p['customer_id'])
    cust = query_one('SELECT * FROM customers WHERE id=?', (cid,))
    if not cust:
        return http.send_json({'error': '客户不存在'}, 404)

    fields = ['company_scale', 'main_products', 'website',
              'comm_best_time', 'comm_style', 'is_urgent_order',
              'own_forwarder', 'custom_level', 'certification', 'trade_terms_detail',
              'risk_notes', 'profile_notes', 'quote_summary', 'website_checked']

    def do_save(c):
        existing = c.execute('SELECT id FROM customer_profiles WHERE customer_id=?', (cid,)).fetchone()
        if existing:
            sets = ','.join(f + '=?' for f in fields)
            vals = [b.get(f, '') for f in fields] + [cid]
            c.execute('UPDATE customer_profiles SET ' + sets + ',updated_at=CURRENT_TIMESTAMP WHERE customer_id=?', vals)
        else:
            cols = ['customer_id'] + fields
            vals = [cid] + [b.get(f, '') for f in fields]
            c.execute('INSERT INTO customer_profiles(' + ','.join(cols) + ') VALUES (' + ','.join('?' * len(cols)) + ')',
                      vals)

    tx(do_save)
    log('customer_profile', cid, 'save', '更新客户画像')
    return {'ok': True}


# ==================== 报价 Excel 导入 ====================

def h_import_excel(p, q, b, http):
    cid = int(p['customer_id'])
    cust = query_one('SELECT * FROM customers WHERE id=?', (cid,))
    if not cust:
        return http.send_json({'error': '客户不存在'}, 404)

    if not openpyxl:
        return http.send_json({'error': '服务器未安装 openpyxl，无法解析 Excel'}, 500)

    b64 = b.get('file_b64', '')
    fname = b.get('file_name', '报价导入.xlsx')
    if not b64:
        return http.send_json({'error': '请上传 Excel 文件'}, 400)

    try:
        raw = base64.b64decode(b64)
    except Exception:
        return http.send_json({'error': '文件编码无效'}, 400)

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as e:
        return http.send_json({'error': '无法解析 Excel 文件：' + str(e)}, 400)

    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        return http.send_json({'error': 'Excel 文件没有数据行（第2行开始）'}, 400)

    items = []
    for r in rows:
        if not r or all(v is None or str(v).strip() == '' for v in r):
            continue
        item = {
            'quote_date': _cell_str(r, 0),
            'quote_no': _cell_str(r, 1),
            'is_closed': _cell_int(r, 2, 0),
            'original_price': _cell_float(r, 3, 0),
            'final_price': _cell_float(r, 4, 0),
            'modification_count': _cell_int(r, 5, 0),
            'is_sample': _cell_int(r, 6, 0),
            'order_type': _cell_str(r, 7),
            'currency': _cell_str(r, 8) or 'USD',
            'exchange_rate': _cell_float(r, 9, None),
            'notes': _cell_str(r, 10),
        }
        items.append(item)

    if not items:
        return http.send_json({'error': '未解析到有效数据行'}, 400)

    # 保存文件和明细
    def do_import(c):
        eid = c.execute(
            'INSERT INTO customer_quote_excels(customer_id,file_name,file_blob,row_count) VALUES (?,?,?,?)',
            (cid, fname, raw, len(items))).lastrowid
        for item in items:
            c.execute(
                'INSERT INTO customer_quote_items(excel_id,quote_date,quote_no,is_closed,'
                'original_price,final_price,modification_count,is_sample,order_type,currency,'
                'exchange_rate,notes,raw_json) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (eid, item['quote_date'], item['quote_no'], item['is_closed'],
                 item['original_price'], item['final_price'], item['modification_count'],
                 item['is_sample'], item['order_type'], item['currency'],
                 item['exchange_rate'], item['notes'], json.dumps(item, ensure_ascii=False)))
        return eid

    eid = tx(do_import)
    log('customer_quote_excel', eid, 'import', '%s: %d rows' % (cust['company'], len(items)))
    return {'ok': True, 'excel_id': eid, 'row_count': len(items)}


def _cell_str(row, idx):
    if idx >= len(row):
        return ''
    v = row[idx]
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    return str(v).strip()


def _cell_int(row, idx, default=0):
    try:
        v = _cell_str(row, idx)
        if v == '':
            return default
        return int(float(v))
    except Exception:
        return default


def _cell_float(row, idx, default=0):
    try:
        v = _cell_str(row, idx)
        if v == '':
            return default
        return float(v)
    except Exception:
        return default


# ==================== 报价统计 ====================

def h_get_stats(p, q, b, http):
    cid = int(p['customer_id'])
    items = query_all(
        'SELECT qi.* FROM customer_quote_items qi '
        'JOIN customer_quote_excels qe ON qe.id=qi.excel_id '
        'WHERE qe.customer_id=? ORDER BY qi.quote_date DESC', (cid,))

    if not items:
        return {
            'total_quotes': 0, 'discount_min': None, 'discount_max': None,
            'avg_modifications': 0, 'sample_count': 0,
            'small_order_count': 0, 'large_order_count': 0,
            'exchange_rate_history': [], 'item_count': 0
        }

    total = len(items)
    closed = [i for i in items if i['is_closed'] == 1 and i['original_price'] and i['original_price'] > 0]
    discounts = [i['final_price'] / i['original_price'] for i in closed if i['final_price']]

    sample_count = sum(1 for i in items if i['is_closed'] == 1 and i['is_sample'] == 1)
    small_count = sum(1 for i in items if i['is_closed'] == 1 and i['order_type'] == '小')
    large_count = sum(1 for i in items if i['is_closed'] == 1 and i['order_type'] == '大')

    mods = [i['modification_count'] for i in items if i['modification_count'] is not None and i['modification_count'] >= 0]

    # 汇率历史
    rates = {}
    for i in items:
        if i['quote_date'] and i['exchange_rate']:
            d = str(i['quote_date'])[:10]
            if d not in rates:
                rates[d] = i['exchange_rate']
    rate_history = sorted([{'date': d, 'rate': r} for d, r in rates.items()], key=lambda x: x['date'])

    return {
        'total_quotes': total,
        'discount_min': round(min(discounts) * 100, 1) if discounts else None,
        'discount_max': round(max(discounts) * 100, 1) if discounts else None,
        'avg_modifications': round(sum(mods) / len(mods), 1) if mods else 0,
        'sample_count': sample_count,
        'small_order_count': small_count,
        'large_order_count': large_count,
        'exchange_rate_history': rate_history,
        'item_count': total
    }


def h_quotation_stats(p, q, b, http):
    """从报价历史(quotations + quotation_items)拉取统计，不依赖Excel导入。"""
    cid = int(p['customer_id'])
    cust = query_one('SELECT * FROM customers WHERE id=?', (cid,))
    if not cust:
        return http.send_json({'error': '客户不存在'}, 404)

    # 该客户的所有报价
    quotes = query_all(
        'SELECT id, quote_no, quote_date, total_usd, status, currency '
        'FROM quotations WHERE customer_id=? ORDER BY quote_date DESC', (cid,))

    total_quotes = len(quotes)
    if total_quotes == 0:
        return {
            'total_quotes': 0, 'total_amount_usd': 0,
            'avg_amount_usd': 0, 'latest_quote_date': None,
            'quotes': [], 'source': 'quotation_history'
        }

    total_amount = sum(q['total_usd'] for q in quotes if q['total_usd'])
    avg_amount = round(total_amount / total_quotes, 2) if total_quotes > 0 else 0

    # 获取所有报价明细的产品统计
    qids = [q['id'] for q in quotes]
    if qids:
        placeholders = ','.join('?' * len(qids))
        items = query_all(
            'SELECT product_name, description, qty, unit_price_usd, amount_usd '
            'FROM quotation_items WHERE quotation_id IN (%s)' % placeholders,
            qids)
        total_items = len(items)
        total_qty = sum(i['qty'] for i in items if i['qty'])
    else:
        items = []
        total_items = 0
        total_qty = 0

    return {
        'total_quotes': total_quotes,
        'total_amount_usd': round(total_amount, 2),
        'avg_amount_usd': avg_amount,
        'total_items': total_items,
        'total_qty': round(total_qty, 2),
        'latest_quote_date': quotes[0]['quote_date'] if quotes else None,
        'recent_quotes': [{'quote_no': q['quote_no'], 'date': q['quote_date'],
                           'total_usd': q['total_usd'], 'status': q['status']}
                          for q in quotes[:10]],
        'source': 'quotation_history'
    }


def h_excel_list(p, q, b, http):
    cid = int(p['customer_id'])
    rows = query_all(
        'SELECT id, file_name, row_count, uploaded_at FROM customer_quote_excels '
        'WHERE customer_id=? ORDER BY uploaded_at DESC', (cid,))
    return rows


def h_delete_excel(p, q, b, http):
    eid = int(p['id'])
    tx(lambda c: c.execute('DELETE FROM customer_quote_excels WHERE id=?', (eid,)))
    log('customer_quote_excel', eid, 'delete', '')
    return {'ok': True}


# ==================== 附件管理 ====================

def h_upload_attachment(p, q, b, http):
    cid = int(p['customer_id'])
    cust = query_one('SELECT * FROM customers WHERE id=?', (cid,))
    if not cust:
        return http.send_json({'error': '客户不存在'}, 404)

    b64 = b.get('file_b64', '')
    fname = b.get('file_name', 'attachment')
    atype = b.get('attachment_type', 'quotation_file')  # quotation_file / whatsapp_file / website_check

    if not b64:
        return http.send_json({'error': '请上传文件'}, 400)

    try:
        raw = base64.b64decode(b64)
    except Exception:
        return http.send_json({'error': '文件编码无效'}, 400)

    def do_upload(c):
        return c.execute(
            'INSERT INTO customer_attachments(customer_id,file_name,attachment_type,file_blob) VALUES (?,?,?,?)',
            (cid, fname, atype, raw)).lastrowid

    aid = tx(do_upload)
    log('customer_attachment', aid, 'upload', '%s: %s' % (cust['company'], fname))
    return {'ok': True, 'attachment_id': aid}


def h_list_attachments(p, q, b, http):
    cid = int(p['customer_id'])
    atype = q.get('type', '')
    sql = 'SELECT id, file_name, attachment_type, uploaded_at FROM customer_attachments WHERE customer_id=?'
    args = [cid]
    if atype:
        sql += ' AND attachment_type=?'
        args.append(atype)
    sql += ' ORDER BY uploaded_at DESC'
    return query_all(sql, args)


def h_delete_attachment(p, q, b, http):
    aid = int(p['id'])
    tx(lambda c: c.execute('DELETE FROM customer_attachments WHERE id=?', (aid,)))
    log('customer_attachment', aid, 'delete', '')
    return {'ok': True}


# ==================== AI 提示词模板 ====================

AI_PROMPTS = {
    'website': {
        'label': '官网背调',
        'desc': '根据客户官网分析公司背景、规模、产品线',
        'template': (
            '请帮我分析以下客户的公司背景：\n'
            '客户名称：{company}\n'
            '官网：{website}\n'
            '国家：{country}\n\n'
            '请从以下维度分析：\n'
            '1. 公司规模与定位（分销商/工程商/终端用户）\n'
            '2. 主营产品线与我司产品的匹配度\n'
            '3. 采购能力与信用评估\n'
            '4. 潜在合作机会与风险点\n\n'
            '--- 以下粘贴官网内容或自行补充 ---\n'
            '{extra}'
        )
    },
    'whatsapp': {
        'label': 'WhatsApp聊天分析',
        'desc': '分析聊天记录提取沟通风格、需求、决策链',
        'template': (
            '请分析以下与客户 {company} 的WhatsApp聊天记录：\n\n'
            '--- 聊天记录 ---\n'
            '{extra}\n'
            '--- 结束 ---\n\n'
            '请提取以下信息：\n'
            '1. 客户沟通风格（直接/委婉/技术型/价格型）\n'
            '2. 最佳联系时间\n'
            '3. 是否急单型客户\n'
            '4. 关键需求与痛点\n'
            '5. 决策链与采购流程\n'
            '6. 价格敏感度\n'
        )
    },
    'quote_analysis': {
        'label': '报价分析',
        'desc': '分析报价数据，总结成交规律与定价策略',
        'template': (
            '请分析以下客户 {company} 的报价数据：\n\n'
            '报价统计：\n'
            '- 近N年报价单数：{total_quotes}\n'
            '- 成交折扣区间：{discount_range}\n'
            '- 平均修改次数：{avg_modifications}\n'
            '- 样品单数：{sample_count}\n'
            '- 大单数：{large_count}，小单数：{small_count}\n\n'
            '--- 附加信息 ---\n'
            '{extra}\n\n'
            '请分析：\n'
            '1. 报价成交规律与定价建议\n'
            '2. 折扣谈判策略\n'
            '3. 样品转化率评估\n'
            '4. 大小单分布特征\n'
            '5. 风险提示\n'
        )
    },
    'comprehensive': {
        'label': '综合画像',
        'desc': '综合所有信息生成完整客户画像',
        'template': (
            '请综合以下信息，为 {company} 生成一份完整的客户画像：\n\n'
            '【基本信息】\n'
            '国家：{country}\n'
            '公司规模：{company_scale}\n'
            '主营产品：{main_products}\n'
            '官网：{website}\n\n'
            '【报价统计】\n'
            '报价单数：{total_quotes}，成交折扣区间：{discount_range}\n'
            '平均修改次数：{avg_modifications}，样品单：{sample_count}\n'
            '大单：{large_count}，小单：{small_count}\n\n'
            '【沟通习惯】\n'
            '最佳联系时间：{comm_best_time}\n'
            '沟通风格：{comm_style}\n'
            '是否急单：{is_urgent}\n\n'
            '【定制物流】\n'
            '自有货代：{own_forwarder}\n'
            '定制等级：{custom_level}\n'
            '认证要求：{certification}\n'
            '贸易术语：{trade_terms}\n\n'
            '【风险备忘】\n'
            '{risk_notes}\n\n'
            '【附加素材】\n'
            '{extra}\n\n'
            '请生成：\n'
            '1. 客户综合画像总结（200字以内）\n'
            '2. 报价策略建议\n'
            '3. 需要关注的3个关键风险点\n'
            '4. 下次沟通建议\n'
        )
    }
}


def h_get_prompts(p, q, b, http):
    return {'prompts': [{'key': k, 'label': v['label'], 'desc': v['desc']} for k, v in AI_PROMPTS.items()]}


# ==================== 导出 ====================

def h_export_profile(p, q, b, http):
    cid = int(p['customer_id'])
    fmt = p['fmt']

    cust = query_one('SELECT * FROM customers WHERE id=?', (cid,))
    if not cust:
        return http.send_json({'error': '客户不存在'}, 404)

    pf = query_one('SELECT * FROM customer_profiles WHERE customer_id=?', (cid,)) or {}

    items = query_all(
        'SELECT qi.* FROM customer_quote_items qi '
        'JOIN customer_quote_excels qe ON qe.id=qi.excel_id '
        'WHERE qe.customer_id=? ORDER BY qi.quote_date DESC', (cid,))

    total = len(items)
    closed = [i for i in items if i['is_closed'] == 1 and i['original_price'] and i['original_price'] > 0]
    discounts = [i['final_price'] / i['original_price'] for i in closed if i['final_price']]
    mods = [i['modification_count'] for i in items if i['modification_count'] is not None and i['modification_count'] >= 0]
    sample_count = sum(1 for i in items if i['is_closed'] == 1 and i['is_sample'] == 1)
    small_count = sum(1 for i in items if i['is_closed'] == 1 and i['order_type'] == '小')
    large_count = sum(1 for i in items if i['is_closed'] == 1 and i['order_type'] == '大')

    disc_range = (
        ('%.1f%% ~ %.1f%%' % (min(discounts) * 100, max(discounts) * 100))
        if discounts else '无数据'
    )
    avg_mod = round(sum(mods) / len(mods), 1) if mods else 0

    if fmt == 'json':
        data = {
            'customer': {k: cust[k] for k in ['company', 'country', 'city', 'customer_type', 'contact', 'email', 'whatsapp_phone', 'currency']},
            'profile': {k: pf.get(k, '') for k in ['company_scale', 'main_products', 'website', 'comm_best_time', 'comm_style', 'own_forwarder', 'custom_level', 'certification', 'trade_terms_detail', 'risk_notes', 'quote_summary']},
            'stats': {
                'total_quotes': total, 'discount_min': round(min(discounts) * 100, 1) if discounts else None,
                'discount_max': round(max(discounts) * 100, 1) if discounts else None,
                'avg_modifications': avg_mod, 'sample_count': sample_count,
                'small_order_count': small_count, 'large_order_count': large_count,
            },
            'items': [{'date': i['quote_date'], 'quote_no': i['quote_no'], 'is_closed': i['is_closed'],
                       'original_price': i['original_price'], 'final_price': i['final_price'],
                       'is_sample': i['is_sample'], 'order_type': i['order_type']} for i in items],
            'exported_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        }
        return http.send_json(data)

    # Markdown export
    md = []
    md.append('# 客户画像：%s\n' % (cust['company'] or '-'))
    md.append('**画像更新时间**：%s\n' % (pf.get('updated_at', '') or datetime.now().strftime('%Y-%m-%d')))

    md.append('## 基础档案\n')
    md.append('| 字段 | 内容 |')
    md.append('|------|------|')
    for label, key in [('客户名称', 'company'), ('国家', 'country'), ('城市', 'city'),
                       ('客户类型', 'customer_type'), ('联系人', 'contact'), ('Email', 'email'),
                       ('电话', 'whatsapp_phone'), ('币种', 'currency')]:
        md.append('| %s | %s |' % (label, cust.get(key, '') or '-'))
    md.append('| 公司规模 | %s |' % (pf.get('company_scale', '') or '-'))
    md.append('| 主营产品 | %s |' % (pf.get('main_products', '') or '-'))
    md.append('| 官网 | %s |' % (pf.get('website', '') or '-'))
    md.append('')

    md.append('## 报价统计\n')
    md.append('| 指标 | 数值 |')
    md.append('|------|------|')
    md.append('| 报价单数 | %d |' % total)
    md.append('| 成交折扣区间 | %s |' % disc_range)
    md.append('| 平均修改次数 | %s |' % avg_mod)
    md.append('| 样品单数 | %d |' % sample_count)
    md.append('| 大单数 | %d |' % large_count)
    md.append('| 小单数 | %d |' % small_count)
    md.append('')

    md.append('## 沟通习惯\n')
    md.append('| 字段 | 内容 |')
    md.append('|------|------|')
    md.append('| 最佳联系时间 | %s |' % (pf.get('comm_best_time', '') or '-'))
    md.append('| 沟通风格 | %s |' % (pf.get('comm_style', '') or '-'))
    md.append('| 是否急单 | %s |' % ('是' if pf.get('is_urgent_order') else '否'))
    md.append('')

    md.append('## 定制物流\n')
    md.append('| 字段 | 内容 |')
    md.append('|------|------|')
    md.append('| 自有货代 | %s |' % (pf.get('own_forwarder', '') or '-'))
    md.append('| 定制等级 | %s |' % (pf.get('custom_level', '') or '-'))
    md.append('| 认证 | %s |' % (pf.get('certification', '') or '-'))
    md.append('| 贸易术语 | %s |' % (pf.get('trade_terms_detail', '') or '-'))
    md.append('')

    md.append('## 风险备忘\n')
    md.append((pf.get('risk_notes', '') or '（无）'))
    md.append('')

    md.append('## 一句话报价摘要\n')
    md.append((pf.get('quote_summary', '') or '（无）'))
    md.append('')

    md.append('---')
    md.append('*导出时间：%s*' % datetime.now().strftime('%Y-%m-%d %H:%M:%S'))

    content = '\n'.join(md)
    b = content.encode('utf-8')
    safe_name = 'customer_profile_%s.md' % (cust['customer_code'] or str(cid))
    return http.send_blob(safe_name, b, 'text/markdown; charset=utf-8')